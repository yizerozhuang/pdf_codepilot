import tkinter as tk
import tkinter.font as tkFont
from tkinter import ttk
import json
import os

from config import CONFIGURATION as conf
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo



def sortby(tree, col, descending):
    """sort tree contents when a column header is clicked on"""
    # grab values to sort

    data = [(tree.set(child, col), child) for child in tree.get_children('')]
    # if the data to be sorted is numeric change to float
    # data =  change_numeric(data)
    # now sort the data in place
    data.sort(reverse=descending)
    for ix, item in enumerate(data):
        tree.move(item[1], '', ix)
    # switch the heading so it will sort in the opposite direction
    tree.heading(col, command=lambda col=col: sortby(tree, col, int(not descending)))

def show_status(data_json):
    state = data_json["State"]
    asana_status_map = {"Pending": "06.Pending",
                        "DWG drawings": "07.DWG drawings",
                        "Done": "08.Done",
                        "Installation": "09.Installation",
                        "Construction Phase": "10.Construction",
                        "Quote Unsuccessful": "11.Quote Unsuccessful"}
    if state["Asana State"] in ["Pending", "DWG drawings", "Done", "Installation", "Construction Phase", "Quote Unsuccessful"]:
        return asana_status_map[state["Asana State"]]
    elif state["Fee Accepted"]:
        return "05.Design"
    elif state["Email to Client"]:
        return "04.Chase Client"
    elif state["Generate Proposal"]:
        return "03.Email To Client"
    elif state["Set Up"]:
        return "02.Preview Fee Proposal"
    else:
        return "01.Set Up"


class SearchBarPage(tk.Frame):
    def __init__(self, app):
        tk.Frame.__init__(self, app.main_frame)
        self.app = app

        self.main_frame = tk.Frame(self)
        self.main_frame.pack(fill=tk.BOTH, expand=1)

        self.main_canvas = tk.Canvas(self.main_frame)
        self.main_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=1, anchor="nw")


        self.main_context_frame = tk.LabelFrame(self.main_canvas, text="Main Context")
        self.main_canvas.create_window((0, 0), window=self.main_context_frame, anchor="nw")
        self.main_context_frame.bind("<Configure>", self.reset_scrollregion)



        self.generate_convert_map()
        self.search_bar()
        self.build_tree()

        self.reset()


    def reset(self):
        # if len(self.app.current_quotation.get())!=0:
        #     save(self.app)
        self.generate_data()
        self.update_data(self.master_project)


    def generate_convert_map(self):
        self.convert_map = {
            "Quote No.": lambda data_json: data_json["Project Info"]["Project"]["Quotation Number"],
            "Proj. No.": lambda data_json: data_json["Project Info"]["Project"]["Project Number"],
            "Project detailed address": lambda data_json: data_json["Project Info"]["Project"]["Project Name"],
            "Shop Name": lambda data_json: data_json["Project Info"]["Project"]["Shop Name"],
            "Scale": lambda data_json: data_json["Project Info"]["Project"]["Proposal Type"],
            "Proj. Type": lambda data_json: data_json["Project Info"]["Project"]["Project Type"],
            "Project Status": show_status,
            "Address To": lambda data_json: data_json["Address_to"],
            "Service Been Engaged": lambda data_json: ", ".join(
                [service["Service"] for service in data_json["Project Info"]["Project"]["Service Type"].values() if
                 service["Include"]]),
            "Apt/Room/Area": lambda data_json: data_json["Project Info"]["Building Features"]["Apt"],
            "Basement/Car Spots": lambda data_json: data_json["Project Info"]["Building Features"]["Basement"],
            "Feature/Notes": lambda data_json: data_json["Project Info"]["Building Features"]["Feature"]
        }
        self.mp_header = list(self.convert_map.keys())


    def generate_data(self):
        database_dir = conf["database_dir"]
        res = []
        # for dir in os.listdir(database_dir):
        #     # print(dir)
        #     if os.path.isdir(os.path.join(database_dir, dir)):
        #         data_dir = os.path.join(database_dir, dir, "data.json")
        #         data_json = json.load(open(data_dir))
        #         res.append(
        #             tuple([self.convert_map[title](data_json) for title in self.mp_header])
        #         )
        mp_dir = os.path.join(database_dir, "mp.json")
        mp_json = json.load(open(mp_dir))
        for project in mp_json.values():
            res.append(tuple([project[key] for key in self.mp_header]))
        # res.sort(key=lambda e: e[0])
        self.master_project = res

    def search_bar(self):
        # search_bar_frame = tk.LabelFrame(self.main_context_frame)
        # search_bar_frame.pack()

        container = ttk.Frame(self.main_context_frame)
        container.pack()


        search_frame = tk.Frame(container)
        search_frame.grid(row=0, column=0, sticky="ew")
        self.entry = tk.Entry(search_frame, font=conf["font"], fg="blue", width=200)
        self.entry.grid(row=0, column=0, sticky="ew")
        tk.Button(search_frame, text="Refresh", bg="Brown", fg="white", command=self.refresh, font=conf["font"]).grid(row=0, column=1)
        # tk.Button(search_frame, text="Export MP Excel", bg="Brown", fg="white", command=self.export_data, font=self.conf["font"]).grid(row=0, column=2)
        self.tree = ttk.Treeview(container, height=35, columns=self.mp_header, show="headings", selectmode="browse")
        treeXScroll = ttk.Scrollbar(container, orient=tk.HORIZONTAL)
        treeXScroll.configure(command=self.main_canvas.xview)
        self.main_canvas.configure(xscrollcommand=treeXScroll.set)
        treeXScroll.grid(row=2, column=0, sticky="ew")
        self.tree.grid(row=1, column=0, sticky='nsew')


        self.entry.bind("<KeyRelease>", self.check)
        self.tree.bind("<<TreeviewSelect>>", self.load_project)

    def build_tree(self):
        for col in self.mp_header:
            self.tree.heading(col, text=col.title(), command=lambda c=col: sortby(self.tree, c, 0))
            # adjust the column's width to the header string
            self.tree.column(col, width=tkFont.Font().measure(col.title()))
    def update_data(self, data):
        self.tree.delete(*self.tree.get_children())
        for item in data:
            self.tree.insert('', 'end', values=item)
            # adjust column's width if necessary to fit each value
            # for ix, val in enumerate(item):
                # col_w = tkFont.Font().measure(val)
                # if self.tree.column(self.mp_header[ix], width=None)<col_w: self.tree.column(self.mp_header[ix], width=col_w)
    def load_project(self, event):

        self.tree.selection()
        selected = self.tree.focus()
        value = self.tree.item(selected, "values")
        # self.app.data["Project Info"]["Project"]["Quotation Number"].set(value[0])

        if len(self.app.quotation_number.get()) != 0:
            self.app.save()

        self.app.set_quotation_number(value)
    def check(self, e, search_string=None):
        if search_string is None:
            typed = self.entry.get().strip()
        else:
            typed = search_string
        if typed == "":
            data = self.master_project
        else:
            data = []
            for item in self.master_project:
                for title in item:
                    if not title is None and typed.lower() in title.lower():
                        data.append(item)
        self.update_data(set(data))
        return data


    def refresh(self):
        self.generate_data()
        self.check(None)
        sortby(self.tree, "Quote No.", True)

    def reset_scrollregion(self, event):
        self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))
